"""
orcamento_parser.py — leitor do Rel_OrcRel (Orçado x Realizado) do Protheus.

Leitor fixo (decisão D1 do 00_GENESIS): o layout é o do relatório do Protheus, e as
colunas são achadas pelo nome no cabeçalho, nunca pela posição. Nada é gravado aqui.

  ler_orcamento(caminho, ano) -> {'orcado', 'realizado', 'centros', 'resumo'}

- orcado / realizado: {(competencia, conta, centro): valor}, com valor na convenção do
  razao.valor do FinHub (crédito menos débito). O relatório já vem assim: receita
  positiva, dedução e despesa negativas. O realizado não é gravado: serve só para
  conferir o realizado que o FinHub calcula do razão.
- conta sem espaço no fim; centro só com dígitos ('01.02.069.11' vira '010206911').
- O ano vem da tela: a coluna ANO do relatório vem vazia ou com o ano errado.
"""

import io
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation

MESES = {"janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
         "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12}

# aba -> {campo: (nomes de coluna aceitos, já normalizados), obrigatório}
COLUNAS = {
    "DADOS": {
        "centro": (("cq2ccusto",), True),
        "conta": (("cq2conta",), True),
        "tipo": (("cq2tpsald",), True),
        "mes": (("mes",), True),
        "orcado": (("orcado",), True),
        "realizado": (("realizado",), True),
    },
    "CTT": {
        "centro": (("ccusto", "cttcusto"), True),
        "descricao": (("descmoeda1", "cttdesc01"), False),
        "servico": (("servico",), False),
        "coordenador": (("coordenador", "coodenador"), False),
        "filial": (("filial",), False),
        "tipo": (("tipo",), False),
    },
}
LIMITE_TEXTO = 120
LIMITE_VALOR = Decimal("9999999999")
MAX_ERROS = 20  # listados de uma vez; a leitura segue até o fim para contar todos
# escada: limites pensados para o Rel_OrcRel de 2026 (~13 MB que expandem ~130 MB); revisar se o relatório crescer
MAX_ARQUIVOS_ZIP, MAX_EXPANDIDO, MAX_RAZAO_EXPANSAO = 2_000, 512 * 1024 * 1024, 100
# escada: teto de linhas da aba DADOS (2026 tem ~78 mil); subir se o relatório de um ano passar disso
MAX_LINHAS = 300_000
_RE_NUMERO = re.compile(r"[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?")
_RE_CENTRO = re.compile(r"[0-9]{1,30}")


class ErroOrcamento(ValueError):
    """Problemas de conteúdo, já com aba e linha. `total` conta todos, mesmo os não listados."""

    def __init__(self, mensagem, total=1):
        super().__init__(mensagem)
        self.total = total


def _norm(texto) -> str:
    """'Desc Moeda 1' -> 'descmoeda1': sem acento, minúsculo, só letra e número."""
    sem_acento = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sem_acento.lower())


def _texto(valor) -> str:
    return " ".join(str(valor).split()) if valor is not None else ""


def normalizar_conta(valor) -> str:
    return _texto(valor)


def normalizar_centro(valor) -> str:
    """Só tira o ponto da máscara. Número inteiro vindo como float (10206911.0) vira '10206911';
    float com fração mantém o ponto e cai na conferência de dígitos."""
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else _texto(valor)
    return _texto(valor).replace(".", "")


def conferir_xlsx(raw: bytes) -> None:
    """Tipo real pelos primeiros bytes (xlsx é zip) e guarda contra arquivo que explode ao abrir."""
    if not raw.startswith(b"PK\x03\x04"):
        raise ValueError("O arquivo não é uma planilha Excel (.xlsx).")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as pacote:
            itens = pacote.infolist()
    except zipfile.BadZipFile as erro:
        raise ValueError("O arquivo não é uma planilha Excel (.xlsx).") from erro
    expandido = sum(item.file_size for item in itens)
    if len(itens) > MAX_ARQUIVOS_ZIP or expandido > MAX_EXPANDIDO or expandido > len(raw) * MAX_RAZAO_EXPANSAO:
        raise ValueError("A planilha é grande demais para abrir com segurança.")


def _valor(bruto, onde: str) -> Decimal:
    if bruto is None or (isinstance(bruto, str) and not bruto.strip()):
        return Decimal(0)
    if isinstance(bruto, bool):
        raise ValueError(f"{onde}: valor inválido.")
    texto = str(bruto).strip()
    if isinstance(bruto, str) and "," in texto:
        if "." in texto and texto.rfind(".") > texto.rfind(","):
            texto = texto.replace(",", "")                     # 1,234.56
        else:
            texto = texto.replace(".", "").replace(",", ".")   # 1.234,56
    if not _RE_NUMERO.fullmatch(texto):
        raise ValueError(f'{onde}: valor "{_texto(bruto)[:20]}" não é número.')
    try:
        numero = Decimal(texto)
    except InvalidOperation:
        raise ValueError(f'{onde}: valor "{_texto(bruto)[:20]}" não é número.') from None
    if not numero.is_finite() or abs(numero) > LIMITE_VALOR:
        raise ValueError(f"{onde}: valor fora do limite.")
    return numero


def _mes(bruto, onde: str) -> int:
    numero = None
    if isinstance(bruto, int) and not isinstance(bruto, bool):
        numero = bruto
    elif isinstance(bruto, float) and bruto.is_integer():
        numero = int(bruto)
    elif isinstance(bruto, str) and re.fullmatch(r"\s*[0-9]{1,2}\s*", bruto):
        numero = int(bruto)
    if numero is not None and 1 <= numero <= 12:
        return numero
    if isinstance(bruto, str) and _norm(bruto) in MESES:
        return MESES[_norm(bruto)]
    raise ValueError(f'{onde}: mês "{_texto(bruto)[:20]}" não reconhecido (janeiro a dezembro ou 1 a 12).')


def _limpo(bruto, onde: str) -> str:
    texto = _texto(bruto)
    if texto and set(texto) <= {"?"}:  # o relatório usa "??????" para coordenador não informado
        return ""
    if len(texto) > LIMITE_TEXTO:
        raise ValueError(f"{onde}: texto com mais de {LIMITE_TEXTO} caracteres.")
    return texto


def _iterar(folha, aba: str, inicio: int, fim=None):
    """(número da linha, valores). Erro do openpyxl no meio da leitura (XML quebrado, célula
    ilegível) vira ValueError com a aba, nunca exceção crua."""
    linhas = folha.iter_rows(min_row=inicio, max_row=fim, values_only=True)
    numero = inicio
    while True:
        try:
            valores = next(linhas)
        except StopIteration:
            return
        except Exception:  # openpyxl levanta ParseError, IndexError, ValueError...
            raise ValueError(f"Aba {aba}: há uma célula ilegível perto da linha {numero}. "
                             "Salve de novo no Excel como .xlsx.") from None
        yield numero, valores
        numero += 1


def _cabecalho(folha, aba: str):
    """(número da linha, {campo: índice}). Cabeçalho = 1ª linha, nas 20 primeiras, que tem todas as obrigatórias."""
    esperado = COLUNAS[aba]
    for numero, valores in _iterar(folha, aba, 1, 20):
        nomes = [_norm(v) for v in valores]
        indice = {}
        for campo, (aceitos, _) in esperado.items():
            achado = next((i for i, n in enumerate(nomes) if n in aceitos), None)
            if achado is not None:
                indice[campo] = achado
        if all(campo in indice for campo, (_, obrig) in esperado.items() if obrig):
            return numero, indice
    faltam = [aceitos[0].upper() for campo, (aceitos, obrig) in esperado.items() if obrig]
    raise ValueError(f"Aba {aba}: cabeçalho não encontrado nas 20 primeiras linhas (colunas esperadas: {', '.join(faltam)}).")


def _linhas(folha, aba: str, inicio: int, indice: dict):
    for numero, valores in _iterar(folha, aba, inicio + 1):
        if numero - inicio > MAX_LINHAS:
            raise ValueError(f"Aba {aba}: mais de {MAX_LINHAS} linhas. Exporte um ano por arquivo.")
        registro = {campo: (valores[i] if i < len(valores) else None) for campo, i in indice.items()}
        if any(_texto(v) for v in registro.values()):
            yield numero, registro


def ler_orcamento(caminho, ano: int) -> dict:
    """Lê e confere o relatório inteiro. Erro de linha não para a leitura: os problemas
    saem juntos (até MAX_ERROS), com aba e linha, para corrigir de uma vez."""
    if isinstance(ano, bool) or not isinstance(ano, int) or not 2000 <= ano <= 2100:
        raise ValueError("Ano do orçamento inválido.")
    from openpyxl import load_workbook
    try:
        with open(caminho, "rb") as arquivo:
            conferir_xlsx(arquivo.read())
    except OSError:
        raise ValueError("Não foi possível ler o arquivo enviado.") from None
    try:
        livro = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as erro:  # openpyxl levanta vários tipos para arquivo corrompido
        raise ValueError("Não foi possível abrir a planilha. Salve de novo no Excel como .xlsx.") from erro

    erros, total_erros = [], 0

    def registrar(mensagem):
        nonlocal total_erros
        total_erros += 1
        if len(erros) < MAX_ERROS:
            erros.append(mensagem)

    try:
        faltam = [aba for aba in COLUNAS if aba not in livro.sheetnames]
        if faltam:
            raise ValueError("A planilha não tem a aba " + ", ".join(faltam)
                             + ". Envie o relatório Orçado x Realizado (Rel_OrcRel) do Protheus.")
        for aba in COLUNAS:
            livro[aba].reset_dimensions()  # a tag <dimension> do arquivo pode mentir e cortar linhas
        centros = {}
        inicio, indice = _cabecalho(livro["CTT"], "CTT")
        for numero, r in _linhas(livro["CTT"], "CTT", inicio, indice):
            onde = f"CTT, linha {numero}"
            try:
                centro = normalizar_centro(r.get("centro"))
                if not _RE_CENTRO.fullmatch(centro):
                    raise ValueError(f'{onde}: centro de custo "{_texto(r.get("centro"))[:20]}" inválido.')
                if centro in centros:
                    raise ValueError(f"{onde}: centro {centro} repetido na aba CTT.")
                centros[centro] = dict(
                    descricao=_limpo(r.get("descricao"), onde) or centro,
                    servico=_limpo(r.get("servico"), onde), coordenador=_limpo(r.get("coordenador"), onde),
                    filial=_limpo(r.get("filial"), onde), tipo=_limpo(r.get("tipo"), onde))
            except ValueError as erro:
                registrar(str(erro))

        orcado, realizado, linhas_orcado = {}, {}, 0
        inicio, indice = _cabecalho(livro["DADOS"], "DADOS")
        for numero, r in _linhas(livro["DADOS"], "DADOS", inicio, indice):
            onde = f"DADOS, linha {numero}"
            try:
                conta, centro = normalizar_conta(r.get("conta")), normalizar_centro(r.get("centro"))
                if not conta:
                    raise ValueError(f"{onde}: conta vazia.")
                if len(conta) > 40:
                    raise ValueError(f"{onde}: código de conta longo demais.")
                if not centro:
                    raise ValueError(f"{onde}: centro de custo vazio.")
                if centro not in centros:
                    raise ValueError(f"{onde}: centro {centro[:20]} não está na aba CTT.")
                chave = (f"{ano}-{_mes(r.get('mes'), onde):02d}", conta, centro)
                tipo = _texto(r.get("tipo"))
                if tipo in ("0", "0.0"):
                    linhas_orcado += 1
                    destino, valor = orcado, _valor(r.get("orcado"), onde)
                elif tipo in ("1", "1.0"):
                    destino, valor = realizado, _valor(r.get("realizado"), onde)
                else:
                    raise ValueError(f'{onde}: tipo de saldo "{tipo[:10]}" desconhecido (0 orçado, 1 realizado).')
                destino[chave] = destino.get(chave, Decimal(0)) + valor
            except ValueError as erro:
                registrar(str(erro))
    finally:
        livro.close()

    if total_erros:
        cabeca = ("1 problema no relatório: " if total_erros == 1 else
                  f"{total_erros} problemas no relatório"
                  + (f" (os {MAX_ERROS} primeiros)" if total_erros > MAX_ERROS else "") + ": ")
        raise ErroOrcamento(cabeca + " · ".join(erros), total_erros)
    if not linhas_orcado:
        raise ValueError("O relatório não tem nenhuma linha de orçado (CQ2_TPSALD 0).")

    orcado = {k: float(v) for k, v in orcado.items()}
    realizado = {k: float(v) for k, v in realizado.items()}
    resumo = dict(ano=ano, centros=len(centros), linhas_orcado=linhas_orcado, chaves_orcado=len(orcado),
                  meses=sorted({k[0] for k in orcado}), total_orcado=round(sum(orcado.values()), 2))
    return dict(orcado=orcado, realizado=realizado, centros=centros, resumo=resumo)


if __name__ == "__main__":  # leitura de um arquivo real, só com contagens na saída
    import sys
    lido = ler_orcamento(sys.argv[1], int(sys.argv[2]))
    print(lido["resumo"] | {"meses": len(lido["resumo"]["meses"]), "chaves_realizado": len(lido["realizado"])})
