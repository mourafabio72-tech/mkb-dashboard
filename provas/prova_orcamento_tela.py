"""Prova da tela Orçado x Realizado (rotas /orcamento, /orcamento/dados, /orcamento/exportar).

Banco temporário com dados inventados.

  1. acesso: anônimo vai para o login; perfil leitura abre a tela e consulta (ver é de todos);
  2. isolamento: usuário do hub restrito à MKB não lê a Gnileb nem o consolidado (403/404), e a tela
     dele não oferece o consolidado;
  3. entrada: empresa, ano, mês, modo, abertura, linha e filtro tortos dão 400 ou 404, nunca 500;
  4. número da tela = número do motor (Fase 3) para os mesmos parâmetros;
  5. exportar Excel: exige CSRF, devolve xlsx com as linhas mandadas, neutraliza célula com fórmula,
     recusa nome e tamanho fora do limite, e grava EXPORT_DADOS no log;
  6. tela: menu sanduíche na ordem do padrão, filtro por coluna, toggles com radiogroup, cartão no hub,
     e a bateria de grep (sem select, CSS na tela, hex, emoji, travessão, popup, innerHTML).
"""
import io
import json
import logging
import os
import re
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
os.environ["DB_PATH"] = str(Path(tempfile.mkdtemp()) / "prova.db")
os.environ.pop("HUB_URL", None)

import app as A  # noqa: E402
from config import MODULO_HUB  # noqa: E402
from ingestion import criar_schema, get_conn, seed_empresas  # noqa: E402
import orcamento_engine as oe  # noqa: E402

falhas = []


def confere(cond, msg):
    print(("ok    " if cond else "FALHA ") + msg)
    if not cond:
        falhas.append(msg)


conn = get_conn()
criar_schema(conn)
seed_empresas(conn)
ROB, FOLHA = "3.1.1.01.01.001", "4.1.1.01.02.015"
conn.execute("INSERT INTO centros_custo (empresa_id, codigo, descricao, servico, coordenador, filial, tipo) VALUES (1,'010206911','Centro A','ASSESSORIA','Ana','RJ','')")
conn.execute("INSERT INTO orcamento_lotes (empresa_id, ano, usuario, linhas, total) VALUES (1, 2026, 'prova', 0, 0)")
conn.executemany("INSERT INTO orcamento (empresa_id, competencia, conta_cod, centro_custo, valor, lote_id) VALUES (1,?,?,?,?,1)",
                 [("2026-01", ROB, "010206911", 1000), ("2026-01", FOLHA, "010206911", -300)])
for n, (eid, conta, valor) in enumerate([(1, ROB, 1100), (1, FOLHA, -320), (2, ROB, 400)]):
    conn.execute("INSERT INTO razao (empresa_id, competencia, data_lanc, conta_cod, documento, centro_custo, valor) VALUES (?,?,?,?,?,?,?)",
                 (eid, "2026-01", "2026-01-10", conta, f"D{n}", "01.02.069.11" if eid == 1 else None, valor))
conn.commit()
conn.close()

cli = A.app.test_client()
TOKEN = "token-da-prova"
capturado = io.StringIO()
logging.getLogger("finhub.orcamento").addHandler(logging.StreamHandler(capturado))


def sessao(usuario=None):
    with cli.session_transaction() as s:
        s.clear()
        if usuario:
            s["usuario_logado"] = usuario
        s["csrf_token"] = TOKEN


ADMIN = {"id": 1, "usuario": "adm", "role": "admin", "nome": "Adm"}
LEITURA = {"id": 2, "usuario": "ler", "role": "leitura", "nome": "Ler"}
HUB_MKB = {"id": 3, "usuario": "cli", "role": "leitura", "nome": "Cli", "origem": "hub", "empresas": ["MKB"], "modulos": [MODULO_HUB]}


def dados(**p):
    base = dict(empresa="mkb", ano="2026", mes="1")
    base.update(p)
    return cli.get("/orcamento/dados", query_string={k: v for k, v in base.items() if v is not None})


# 1. acesso
sessao()
confere(cli.get("/orcamento").status_code == 302 and dados().status_code == 302, "acesso: anônimo vai para o login")
sessao(LEITURA)
r = cli.get("/orcamento")
confere(r.status_code == 200 and 'id="orcTela"' in r.get_data(as_text=True), "acesso: perfil leitura abre a tela")
confere(dados().status_code == 200, "acesso: perfil leitura consulta")

# 2. isolamento
sessao(HUB_MKB)
confere(dados().status_code == 200, "isolamento: cliente do hub lê a MKB")
confere(dados(empresa="gnileb").status_code in (403, 404), "isolamento: cliente da MKB não lê a Gnileb")
confere(dados(empresa="consolidado").status_code == 404, "isolamento: cliente da MKB não lê o consolidado")
html = cli.get("/orcamento").get_data(as_text=True)
confere('data-value="consolidado"' not in html and 'data-value="gnileb"' not in html, "isolamento: a tela dele não oferece Gnileb nem consolidado")
confere("ACESSO_NEGADO_IDOR" in capturado.getvalue(), "isolamento: tentativa registrada como ACESSO_NEGADO_IDOR")

# 3. entrada torta
sessao(ADMIN)
casos = [dict(empresa="xyz"), dict(empresa=""), dict(ano="20x6"), dict(ano="1999"), dict(mes="13"), dict(mes="0"),
         dict(modo="ano"), dict(abrir_por="conta"), dict(linha="DROP TABLE"), dict(servico="x" * 200),
         dict(ano="２０２６")]
for c in casos:
    r = dados(**c)
    confere(r.status_code in (400, 404) and "erro" in (r.get_json() or {}), f"entrada: {c} = {r.status_code}")

# 4. número da tela = motor
conn = get_conn()
for p in (dict(), dict(abrir_por="centro"), dict(empresa="consolidado"), dict(abrir_por="servico", linha="ROB"),
          dict(modo="acumulado"), dict(servico="ASSESSORIA")):
    j = dados(**p).get_json()["dados"]
    filtros = {"servico": [p["servico"]]} if "servico" in p else {}
    esperado = oe.calcular(conn, p.get("empresa", "mkb"), 2026, 1, p.get("modo", "mes"), p.get("abrir_por", "linha"), filtros, p.get("linha"))
    confere(j["total"] == esperado["total"] and len(j["linhas"]) == len(esperado["linhas"]), f"número: tela = motor em {p or 'padrão'}")
j = dados().get_json()["dados"]
confere(j["total"]["orcado"] == 700 and j["total"]["realizado"] == 780, "número: total MKB 700 x 780 (conta à mão)")
confere(dados(empresa="consolidado").get_json()["dados"]["total"]["realizado"] == 1180, "número: consolidado 1180")
conn.close()

# 4b. rota da árvore (Fase 6)
def arvore_get(**p):
    base = dict(empresa="mkb", ano="2026", de="1", ate="1")
    base.update(p)
    return cli.get("/orcamento/arvore", query_string={k: v for k, v in base.items() if v is not None})
sessao(LEITURA)
r = arvore_get()
conn = get_conn()
confere(r.status_code == 200 and r.get_json()["dados"]["total"] == oe.arvore(conn, "mkb", 2026, 1, 1)["total"], "árvore: leitura consulta e tela = motor")
conn.close()
confere(r.get_json()["dados"]["total"]["acumulado"] == [700, 780], "árvore: total MKB 700 x 780 (conta à mão)")
for c in (dict(de="2", ate="1"), dict(de="0"), dict(ate="13"), dict(de="x"), dict(empresa="xyz")):
    confere(arvore_get(**c).status_code in (400, 404), f"árvore: entrada {c} recusada")
sessao(HUB_MKB)
confere(arvore_get(empresa="gnileb").status_code in (403, 404) and arvore_get(empresa="consolidado").status_code == 404,
        "árvore: cliente da MKB não lê Gnileb nem consolidado")
sessao(ADMIN)

# 5. exportar Excel
corpo = {"nome": "orcamento_tabela_2026-10-20", "cabecalho": ["Linha da DRE", "Orçado"], "linhas": [["=1+1", 10.5], ["Receita", -3]]}
r = cli.post("/orcamento/exportar", json=corpo)
confere(r.status_code == 400, "exportar: sem token CSRF = 400")
r = cli.post("/orcamento/exportar", json=corpo, headers={"X-CSRF-Token": TOKEN})
confere(r.status_code == 200 and r.data[:4] == b"PK\x03\x04", "exportar: devolve xlsx")
from openpyxl import load_workbook  # noqa: E402
folha = load_workbook(io.BytesIO(r.data)).active
confere([c.value for c in folha[2]] == ["'=1+1", 10.5] and folha.max_row == 3, "exportar: linhas mandadas, fórmula neutralizada")
confere("EXPORT_DADOS" in capturado.getvalue(), "exportar: EXPORT_DADOS no log")
for ruim in (dict(corpo, nome="../../x"), dict(corpo, linhas=[["a"]]), dict(corpo, linhas=[["a", 1]] * 5001),
             dict(corpo, cabecalho=["c"] * 61)):
    confere(cli.post("/orcamento/exportar", json=ruim, headers={"X-CSRF-Token": TOKEN}).status_code == 400,
            f"exportar: pedido fora do limite = 400 ({len(ruim['linhas'])} linhas)")
r = cli.post("/orcamento/exportar", json=dict(corpo, cabecalho=["a\x07", "b"], linhas=[["c\x01\x0b", ""]]), headers={"X-CSRF-Token": TOKEN})
confere(r.status_code == 200 and [c.value for c in load_workbook(io.BytesIO(r.data)).active[2]] == ["c", None], "exportar: caractere de controle tirado, vazio continua vazio")
sessao(HUB_MKB)
conn = get_conn()
conn.execute("INSERT INTO razao (empresa_id, competencia, data_lanc, conta_cod, documento, valor) VALUES (2,'2031-05','2031-05-10',?,'Z1',1)", (ROB,))
conn.commit(); conn.close()
confere("2031" not in cli.get("/orcamento").get_data(as_text=True), "isolamento: ano que só a Gnileb tem não aparece para o cliente da MKB")
sessao(ADMIN)
confere(cli.post("/orcamento/exportar", data="isto não é json", headers={"X-CSRF-Token": TOKEN},
                 content_type="application/json").status_code == 400, "exportar: corpo que não é JSON = 400")

# 6. tela e bateria de grep
tpl = (RAIZ / "templates" / "orcamento.html").read_text(encoding="utf-8")
js = (RAIZ / "static" / "orcamento_tela.js").read_text(encoding="utf-8")
css = (RAIZ / "static" / "orcamento.css").read_text(encoding="utf-8")
menu_tabela = tpl[tpl.index('id="orcMenu"'):tpl.index('id="orcPopOrdenar"')]
menu_arvore = tpl[tpl.index('id="orcArvMenu"'):tpl.index('id="orcArvore"')]
confere(re.findall(r'data-acao="([a-z]+)"', menu_tabela) == ["filtro", "ordenar", "colunas", "csv", "excel"],
        "tela: menu da tabela na ordem Filtro, Ordenar, Colunas, CSV, Excel")
confere(re.findall(r'data-acao="([a-z]+)"', menu_arvore) == ["filtro", "csv", "excel", "expandir", "recolher"],
        "tela: menu da árvore com Filtro, CSV, Excel, Expandir e Recolher (sem Ordenar e Colunas: grade hierárquica)")
confere("ph-list" in tpl and "ph-dots-three" not in tpl + js and not re.search(r'id="[a-zA-Z]*[Tt]oolbar"', tpl), "tela: sanduíche dentro do card, sem toolbar solta")
confere(tpl.count('placeholder="filtrar"') == 1 and "addEventListener('input', corpo)" in js and "inp.placeholder = 'filtrar'" in js,
        "tela: filtro por coluna a cada caractere, na tabela e na árvore")
confere(tpl.count('role="radiogroup"') == 2, "tela: os dois toggles com radiogroup")
confere(tpl.count('class="cv-sw"') == 2 and re.search(r"\.cv-sw input\s*\{\s*display:\s*none", css), "tela: liga/desliga ÷ 1.000 e centavos no padrão cv-sw")
confere(not re.search(r"\.tabela-wrap\s*\{[^}]*overflow\s*:\s*hidden", css), "tela: .tabela-wrap sem overflow hidden")
confere("tabular-nums" in css and "min-height: 25px" in css, "tela: KPI com tabular-nums e lbl de 25px")
confere("<select" not in tpl and "<style" not in tpl and 'style="' not in tpl, "tela: sem select, sem CSS na tela")
confere(not re.search(r"#[0-9a-fA-F]{3,6}\b", tpl + css) and not re.search(r"rgba?\(", css), "tela: sem hex nem rgba")
confere(not re.search("[\U0001F300-\U0001FAFF☀-➿]", tpl + js) and "—" not in tpl + js + css, "tela: sem emoji nem travessão")
confere(not re.search(r"\b(alert|confirm|prompt)\(", js) and "innerHTML" not in js, "tela: sem popup nativo nem innerHTML")
confere(not re.search(r"Carregando\.\.\.|Aguarde\.\.\.|Processando\.\.\.", tpl + js), "tela: loading sem texto genérico")
confere("toISOString" not in js, "tela: data local, sem toISOString")
sessao(ADMIN)
inicio = cli.get("/").get_data(as_text=True)
confere("/orcamento?competencia=" in inicio and ("Orçado x Realizado" in inicio or "<span>Orçamento</span>" in inicio),
        "tela: cartão no hub do dashboard (grande ou na faixa compacta)")

print()
print(f"{len(falhas)} falha(s)" if falhas else "prova_orcamento_tela: tudo ok")
sys.exit(1 if falhas else 0)
