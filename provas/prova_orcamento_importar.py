"""Prova da tela Importar orçamento (orcamento_rotas.py), com o cliente de teste do Flask.

Banco temporário e planilha gerada aqui, com dados inventados.

  1. acesso: anônimo e usuário `leitura` não entram (nem na tela nem no POST) e nada é gravado;
  2. CSRF: POST sem token dá 400;
  3. entrada: empresa fora da lista dá 404 e loga ACESSO_NEGADO_IDOR; ano torto dá 400;
     sem arquivo dá 400; .exe renomeado para .xlsx dá 400; corpo acima de 30 MB dá 413;
     planilha com erro de linha dá 400 com a linha na mensagem;
  4. prévia: devolve os números e NÃO grava nada;
  5. confirmar: grava orçado, lote e centros; importar de novo o mesmo ano não duplica e
     informa quantas linhas saíram; outro ano e outra empresa ficam intactos;
  6. log: ORCAMENTO_IMPORTADO em JSON com empresa, ano e contagens, sem conta, centro nem valor de linha;
  7. tela: sem <select>, sem CSS na tela, sem hex, sem emoji, sem travessão, sem popup nativo.

Uso:  uv run --python 3.12 --with-requirements requirements.txt python provas/prova_orcamento_importar.py
"""
import io
import json
import logging
import os
import re
import sqlite3
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
TMP = Path(tempfile.mkdtemp())
os.environ["DB_PATH"] = str(TMP / "prova.db")
os.environ.pop("HUB_URL", None)

from openpyxl import Workbook  # noqa: E402

import app as A  # noqa: E402
from ingestion import criar_schema, get_conn, seed_empresas  # noqa: E402

conn = get_conn()
criar_schema(conn)
seed_empresas(conn)
conn.close()

falhas = []


def confere(cond, msg):
    print(("ok    " if cond else "FALHA ") + msg)
    if not cond:
        falhas.append(msg)


def planilha(linhas_dados, centro_conta="CONTA-SECRETA-9"):
    livro = Workbook()
    livro.remove(livro.active)
    ctt = livro.create_sheet("CTT")
    ctt.append(["CTT"]); ctt.append([])
    ctt.append(["C Custo", "Desc Moeda 1", "SERVIÇO", "COORDENADOR", "FILIAL", "TIPO"])
    ctt.append(["01.02.069.11", "Centro A", "ASSESSORIA", "Ana", "RJ", "D"])
    ctt.append(["010209804", "Centro B", "FACILITIES", "??????", "SP", "I"])
    dados = livro.create_sheet("DADOS")
    dados.append(["CQ2_CCUSTO", "CQ2_CONTA", "CQ2_TPSALD", "MES", "ORCADO", "REALIZADO"])
    for linha in linhas_dados:
        dados.append(linha)
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


BOA = planilha([
    ["010206911", "3.1.1.01.01.001   ", "0", "Janeiro", 1000, " "],
    ["010206911", "3.1.1.01.01.001", "0", "Fevereiro", 1100, " "],
    ["010209804", "4.1.1.01.01.001", "0", "Janeiro", -400, " "],
    ["010209804", "9.9.9.99.99.999", "0", "Março", -50, " "],     # conta sem linha na DRE
    ["010209804", "4.1.1.01.01.001", "1", "Janeiro", " ", -380],
])
cli = A.app.test_client()
TOKEN = "token-da-prova"


def sessao(role):
    with cli.session_transaction() as s:
        s.clear()
        if role:
            s["usuario_logado"] = {"id": 7, "usuario": "prova", "role": role, "nome": "Prova"}
        s["csrf_token"] = TOKEN


def post(rota, empresa="mkb", ano="2026", raw=BOA, nome="Rel_OrcRel_2026.xlsx", token=TOKEN):
    dados = {"empresa": empresa, "ano": ano}
    if raw is not None:
        dados["arquivo"] = (io.BytesIO(raw), nome)
    headers = {"X-CSRF-Token": token} if token else {}
    return cli.post(rota, data=dados, headers=headers, content_type="multipart/form-data")


def contar(sql, *p):
    c = sqlite3.connect(os.environ["DB_PATH"])
    try:
        return c.execute(sql, p).fetchone()[0]
    finally:
        c.close()


PREVIA, CONFIRMAR, TELA = "/orcamento/importar/previa", "/orcamento/importar/confirmar", "/orcamento/importar"

capturado = io.StringIO()
logging.getLogger("finhub.orcamento").addHandler(logging.StreamHandler(capturado))

# 1. acesso
sessao(None)
confere(cli.get(TELA).status_code == 302, "acesso: anônimo na tela vai para o login")
confere(post(CONFIRMAR).status_code == 302, "acesso: anônimo no confirmar vai para o login")
sessao("leitura")
confere(cli.get(TELA).status_code == 302, "acesso: perfil leitura não abre a tela")
confere(post(CONFIRMAR).status_code == 302, "acesso: perfil leitura não grava")
confere(contar("SELECT COUNT(*) FROM orcamento") == 0, "acesso: nada gravado por quem não é admin")

# 2. CSRF
sessao("admin")
confere(post(PREVIA, token=None).status_code == 400, "csrf: POST sem token dá 400")
confere(post(CONFIRMAR, token="outro").status_code == 400, "csrf: token errado dá 400")

# 3. entrada
r = post(PREVIA, empresa="xyz")
confere(r.status_code == 404 and "ACESSO_NEGADO_IDOR" in capturado.getvalue(), "entrada: empresa fora da lista = 404 com ACESSO_NEGADO_IDOR")
confere(post(PREVIA, empresa="").status_code == 404, "entrada: empresa vazia = 404")
for ano in ("26", "abcd", "1999", "2026;", "２０２６"):
    confere(post(PREVIA, ano=ano).status_code == 400, f"entrada: ano {ano!r} = 400")
confere(post(PREVIA, raw=None).status_code == 400, "entrada: sem arquivo = 400")
r = post(PREVIA, raw=b"MZ\x90\x00 programa", nome="orcamento.xlsx")
confere(r.status_code == 400 and "não é uma planilha" in r.get_json()["erro"], "entrada: .exe renomeado = 400")
r = cli.post(PREVIA, data=b"x" * (31 * 1024 * 1024), headers={"X-CSRF-Token": TOKEN},
             content_type="application/octet-stream")
confere(r.status_code == 413, "entrada: corpo acima de 30 MB = 413")
r = post(PREVIA, raw=planilha([["010206911", "3.1", "0", "Trezembro", 5, ""]]))
confere(r.status_code == 400 and "DADOS, linha 2" in r.get_json()["erro"], "entrada: erro de linha volta com a linha")
r = cli.post(PREVIA, input_stream=io.BytesIO(b"x" * 10), headers={"X-CSRF-Token": TOKEN, "Transfer-Encoding": "chunked"},
             content_type="multipart/form-data; boundary=x")
confere(r.status_code == 411, "entrada: corpo sem tamanho declarado = 411")
sobras = [f for f in os.listdir(tempfile.gettempdir()) if f.startswith("orc_")]
confere(not sobras, "entrada: nenhum temporário orc_ deixado para trás")
confere(contar("SELECT COUNT(*) FROM orcamento") == 0, "entrada: nada gravado nas recusas")

# 4. prévia
r = post(PREVIA)
p = (r.get_json() or {}).get("previa", {})
confere(r.status_code == 200 and p.get("linhas") == 4 and p.get("centros") == 2, "prévia: 4 linhas e 2 centros")
confere(p.get("total") == 1650.0 and [m["mes"] for m in p.get("meses", [])] == ["Jan", "Fev", "Mar"], "prévia: total e meses")
confere(p.get("qtd_contas_sem_linha") == 1 and p["contas_sem_linha"][0]["conta"] == "9.9.9.99.99.999", "prévia: conta sem linha na DRE listada")
confere(p.get("qtd_centros_novos") == 2 and p.get("linhas_substituidas") == 0, "prévia: centros novos e nada a substituir")
confere(contar("SELECT COUNT(*) FROM orcamento") == 0 and contar("SELECT COUNT(*) FROM orcamento_lotes") == 0, "prévia: não grava")

# 5. confirmar
r = post(CONFIRMAR)
j = r.get_json() or {}
confere(r.status_code == 200 and j.get("linhas") == 4 and j.get("substituidas") == 0, "confirmar: 4 linhas gravadas")
confere(contar("SELECT COUNT(*) FROM orcamento WHERE empresa_id = 1") == 4, "confirmar: 4 linhas da MKB no banco")
confere(contar("SELECT valor FROM orcamento WHERE conta_cod = '3.1.1.01.01.001' AND centro_custo = '010206911' AND competencia = '2026-01'") == 1000,
        "confirmar: chave gravada com conta sem espaço e centro sem ponto")
confere(contar("SELECT COUNT(*) FROM centros_custo WHERE empresa_id = 1") == 2, "confirmar: 2 centros gravados")
confere(contar("SELECT coordenador FROM centros_custo WHERE codigo = '010209804'") == "", "confirmar: '??????' gravado vazio")
confere(contar("SELECT COUNT(*) FROM orcamento_lotes") == 1, "confirmar: 1 lote")
post(CONFIRMAR, ano="2025")
post(CONFIRMAR, empresa="gnileb")
r = post(CONFIRMAR)
j = r.get_json() or {}
confere(j.get("substituidas") == 4 and contar("SELECT COUNT(*) FROM orcamento WHERE empresa_id = 1 AND competencia LIKE '2026-%'") == 4,
        "confirmar: reimportar o mesmo ano não duplica e conta as que saíram")
confere(contar("SELECT COUNT(*) FROM orcamento WHERE empresa_id = 1 AND competencia LIKE '2025-%'") == 4, "confirmar: outro ano intacto")
confere(contar("SELECT COUNT(*) FROM orcamento WHERE empresa_id = 2") == 4, "confirmar: outra empresa intacta")
p = post(PREVIA).get_json()["previa"]
confere(p["linhas_substituidas"] == 4 and p["qtd_centros_novos"] == 0, "prévia: avisa o que vai substituir")

# 6. log
eventos = [json.loads(l) for l in capturado.getvalue().splitlines() if l.startswith("{")]
imp = [e for e in eventos if e["event"] == "ORCAMENTO_IMPORTADO"]
confere(len(imp) == 4 and imp[0]["empresa"] == "mkb" and imp[0]["ano"] == 2026 and imp[0]["linhas"] == 4,
        "log: ORCAMENTO_IMPORTADO com empresa, ano e linhas")
confere(all({"timestamp", "level", "user_id", "ip", "path", "request_id"} <= set(e) for e in imp), "log: campos fixos presentes, com request_id")
confere(len({e["request_id"] for e in imp}) == len(imp), "log: request_id diferente por pedido")
texto_log = capturado.getvalue()
confere("3.1.1.01.01.001" not in texto_log and "010206911" not in texto_log and "Centro A" not in texto_log,
        "log: sem conta, centro nem descrição")

# 7. tela
sessao("admin")
html = cli.get(TELA).get_data(as_text=True)
confere('class="selectbusca"' in html and 'class="im-drop"' in html, "tela: SelectBusca e área de arquivo presentes")
tpl = (RAIZ / "templates" / "orcamento_importar.html").read_text(encoding="utf-8")
js = "".join((RAIZ / "static" / f).read_text(encoding="utf-8") for f in ("orcamento.js", "selectbusca.js"))
css = (RAIZ / "static" / "orcamento.css").read_text(encoding="utf-8")
confere("<select" not in tpl, "tela: sem <select> nativo")
confere("<style" not in tpl and 'style="' not in tpl, "tela: sem CSS na tela")
confere(not re.search(r"#[0-9a-fA-F]{3,6}\b", tpl + css), "tela: sem hex no template nem no CSS do módulo")
confere(not re.search(r"rgba?\(", css), "tela: sem rgba solto no CSS do módulo")
confere(not re.search("[\U0001F300-\U0001FAFF☀-➿]", tpl + js), "tela: sem emoji")
confere("—" not in tpl + js + css, "tela: sem travessão")
confere(not re.search(r"\b(alert|confirm|prompt)\(", js), "tela: sem popup nativo")
confere(not re.search(r"Carregando\.\.\.|Aguarde\.\.\.|Processando\.\.\.", tpl + js), "tela: loading sem texto genérico")
confere(tpl.index('class="im-msg"') < tpl.index('id="orcLer"'), "tela: mensagem antes do botão")
confere("innerHTML" not in js, "tela: JS sem innerHTML")

print()
print(f"{len(falhas)} falha(s)" if falhas else "prova_orcamento_importar: tudo ok")
sys.exit(1 if falhas else 0)
