"""Prova do leitor do Rel_OrcRel (orcamento_parser.py) e das tabelas do orçamento.

A planilha é gerada aqui, com dados inventados: nenhum dado de cliente entra no git.

  1. schema: criar_schema duas vezes no mesmo banco cria orcamento, orcamento_lotes e
     centros_custo, com a UNIQUE de empresa + competência + conta + centro, e não quebra;
  2. leitura: conta sem espaço no fim, centro sem ponto (na CTT e na DADOS), mês por
     extenso e por número, sinal mantido (crédito menos débito), só o tipo 0 entra no
     orçado, chave repetida soma, '??????' vira vazio, coluna fora de ordem é achada pelo nome;
  3. erros com aba e linha, todos juntos: conta vazia, centro vazio, centro fora da CTT,
     mês fora de 1 a 12, valor que não é número, tipo de saldo desconhecido;
  4. aba faltando, coluna obrigatória faltando, ano inválido e arquivo que não é xlsx
     dão erro claro.

Uso:  uv run --python 3.12 --with-requirements requirements.txt python provas/prova_orcamento_leitor.py
"""
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
TMP = Path(tempfile.mkdtemp())
os.environ["DB_PATH"] = str(TMP / "prova.db")

from openpyxl import Workbook  # noqa: E402

import orcamento_parser as op  # noqa: E402

CAB_DADOS = ["CQ2_CCUSTO", "CTT_DESC01", "CQ2_CONTA", "DEBITO", "CREDITO", "CQ2_TPSALD", "MES",
             "ORCADO", "REALIZADO", "ANO"]
CAB_CTT = ["C Custo", "Desc Moeda 1", "SERVIÇO", "COORDENADOR", "FILIAL", "TIPO"]
falhas = []


def confere(cond, msg):
    print(("ok    " if cond else "FALHA ") + msg)
    if not cond:
        falhas.append(msg)


def planilha(nome, dados, ctt=None, abas=("CTT", "DADOS"), cab_dados=CAB_DADOS):
    livro = Workbook()
    livro.remove(livro.active)
    if "CTT" in abas:
        folha = livro.create_sheet("CTT")
        folha.append(["CTT"])            # o relatório real tem o cabeçalho da CTT na linha 3
        folha.append([])
        folha.append(CAB_CTT)
        for linha in ctt or [["01.02.069.11", "Centro A", "ASSESSORIA", "Ana", "RJ", "D"],
                             ["010209804", "Centro B", "FACILITIES", "??????", "SP", "I"]]:
            folha.append(linha)
    if "DADOS" in abas:
        folha = livro.create_sheet("DADOS")
        folha.append(cab_dados)
        for linha in dados:
            folha.append(linha)
    caminho = TMP / nome
    livro.save(caminho)
    return caminho


def erro_de(funcao):
    try:
        funcao()
    except ValueError as erro:
        return erro
    return None


# 1. schema
import ingestion  # noqa: E402

conn = sqlite3.connect(os.environ["DB_PATH"])
ingestion.criar_schema(conn)
ingestion.criar_schema(conn)
tabelas = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
confere({"orcamento", "orcamento_lotes", "centros_custo"} <= tabelas, "schema: as 3 tabelas existem depois de 2 boots")
conn.execute("INSERT INTO orcamento_lotes (empresa_id, ano, usuario, linhas, total) VALUES (1, 2026, 'prova', 1, 10)")
conn.execute("INSERT INTO orcamento (empresa_id, competencia, conta_cod, centro_custo, valor, lote_id) VALUES (1,'2026-01','3.1','01',10,1)")
try:
    conn.execute("INSERT INTO orcamento (empresa_id, competencia, conta_cod, centro_custo, valor, lote_id) VALUES (1,'2026-01','3.1','01',5,1)")
    unica = False
except sqlite3.IntegrityError:
    unica = True
confere(unica, "schema: UNIQUE barra a mesma empresa + mês + conta + centro")
conn.execute("INSERT INTO orcamento (empresa_id, competencia, conta_cod, centro_custo, valor, lote_id) VALUES (2,'2026-01','3.1','01',5,1)")
confere(True, "schema: a mesma chave em outra empresa entra")
conn.close()

# 2. leitura
boa = planilha("boa.xlsx", [
    ["01.02.069.11", "Centro A", "3.1.1.01.01.001   ", 0, 0, "0", "Janeiro", 1000.5, " ", None],
    ["010206911", "Centro A", "3.1.1.01.01.001", 0, 0, "0", "janeiro", 499.5, " ", 2025],     # mesma chave: soma
    ["010206911", "Centro A", "3.1.3.01.02.001   ", 0, 0, 0, 3, -87.57, " ", None],             # dedução negativa, mês número
    ["010209804", "Centro B", "4.1.1.01.01.001", 0, 0, "0", "Março", "-1.234,56", " ", None],    # texto em formato brasileiro
    ["010209804", "Centro B", "4.1.1.01.01.001", 0, 0, "1", "Março", " ", -1200, None],          # realizado: fora do orçado
])
lido = op.ler_orcamento(boa, 2026)
o = lido["orcado"]
confere(set(o) == {("2026-01", "3.1.1.01.01.001", "010206911"), ("2026-03", "3.1.3.01.02.001", "010206911"),
                   ("2026-03", "4.1.1.01.01.001", "010209804")}, "leitura: chaves com conta sem espaço, centro sem ponto e ano da tela")
confere(o[("2026-01", "3.1.1.01.01.001", "010206911")] == 1500.0, "leitura: chave repetida soma (1000,50 + 499,50)")
confere(o[("2026-03", "3.1.3.01.02.001", "010206911")] == -87.57, "leitura: sinal mantido (dedução negativa)")
confere(o[("2026-03", "4.1.1.01.01.001", "010209804")] == -1234.56, "leitura: '-1.234,56' vira -1234,56")
confere(lido["realizado"] == {("2026-03", "4.1.1.01.01.001", "010209804"): -1200.0}, "leitura: tipo 1 vai só para o realizado de conferência")
confere(lido["resumo"]["linhas_orcado"] == 4 and lido["resumo"]["chaves_orcado"] == 3, "leitura: 4 linhas de orçado, 3 chaves")
confere(lido["resumo"]["total_orcado"] == round(1500 - 87.57 - 1234.56, 2), "leitura: total do orçado")
confere(set(lido["centros"]) == {"010206911", "010209804"} and lido["centros"]["010209804"]["coordenador"] == "",
        "leitura: CTT sem ponto e '??????' vira coordenador vazio")
confere(lido["centros"]["010206911"] == dict(descricao="Centro A", servico="ASSESSORIA", coordenador="Ana", filial="RJ", tipo="D"),
        "leitura: campos do centro")
fora_de_ordem = planilha("ordem.xlsx", [["3.1.1.01.01.001", "0", "Maio", 10, "", "010206911"]],
                         cab_dados=["CQ2_CONTA", "CQ2_TPSALD", "MES", "ORCADO", "REALIZADO", "CQ2_CCUSTO"])
confere(op.ler_orcamento(fora_de_ordem, 2026)["orcado"] == {("2026-05", "3.1.1.01.01.001", "010206911"): 10.0},
        "leitura: colunas em outra ordem são achadas pelo nome")

# 3. erros com aba e linha
ruim = planilha("ruim.xlsx", [
    ["010206911", "x", "", 0, 0, "0", "Janeiro", 1, "", None],           # linha 2: conta vazia
    ["", "x", "3.1", 0, 0, "0", "Janeiro", 1, "", None],                 # linha 3: centro vazio
    ["999", "x", "3.1", 0, 0, "0", "Janeiro", 1, "", None],              # linha 4: centro fora da CTT
    ["010206911", "x", "3.1", 0, 0, "0", "Trezembro", 1, "", None],      # linha 5: mês
    ["010206911", "x", "3.1", 0, 0, "0", 13, 1, "", None],               # linha 6: mês 13
    ["010206911", "x", "3.1", 0, 0, "0", "Janeiro", "abc", "", None],    # linha 7: valor
    ["010206911", "x", "3.1", 0, 0, "7", "Janeiro", 1, "", None],        # linha 8: tipo
])
erro = erro_de(lambda: op.ler_orcamento(ruim, 2026))
texto = str(erro)
confere(isinstance(erro, op.ErroOrcamento) and erro.total == 7, "erros: os 7 problemas contados de uma vez")
for linha, trecho in ((2, "conta vazia"), (3, "centro de custo vazio"), (4, "não está na aba CTT"), (5, "Trezembro"),
                      (6, '"13"'), (7, "não é número"), (8, "tipo de saldo")):
    confere(f"DADOS, linha {linha}: " in texto and trecho in texto, f"erros: linha {linha} com '{trecho}'")

grande = planilha("grande.xlsx", [["010206911", "x", "3.1", 0, 0, "0", "Janeiro", 1e11, "", None]])
confere("DADOS, linha 2: valor fora do limite" in str(erro_de(lambda: op.ler_orcamento(grande, 2026))),
        "erros: valor acima de 9.999.999.999 recusado com a linha")
ctt_ruim = planilha("ctt_ruim.xlsx", [["010206911", "x", "3.1", 0, 0, "0", "Janeiro", 1, "", None]],
                    ctt=[["010206911", "A", "", "", "", ""], ["CC-ABC", "B", "", "", "", ""]])
confere('CTT, linha 5: centro de custo "CC-ABC" inválido' in str(erro_de(lambda: op.ler_orcamento(ctt_ruim, 2026))),
        "erros: centro com letra na CTT recusado com a linha")

# 3b. furos achados pelo verificador (2026-10-20): mês, centro e valor que viravam outra coisa em silêncio
mes_float = planilha("mes_float.xlsx", [["010206911", "x", "3.1", 0, 0, "0", 1.0, 5, "", None]])
confere(set(op.ler_orcamento(mes_float, 2026)["orcado"]) == {("2026-01", "3.1", "010206911")}, "furo: mês 1.0 é janeiro, não outubro")
for mes in (1.2, 0.5, "1-2", "1.0"):
    p = planilha("mes_ruim.xlsx", [["010206911", "x", "3.1", 0, 0, "0", mes, 5, "", None]])
    confere("não reconhecido" in str(erro_de(lambda: op.ler_orcamento(p, 2026))), f"furo: mês {mes!r} recusado")
centro_float = planilha("centro_float.xlsx", [[10206911.0, "x", "3.1", 0, 0, "0", "Maio", 5, "", None]],
                        ctt=[["102069110", "A", "", "", "", ""]])
confere("centro 10206911 não está na aba CTT" in str(erro_de(lambda: op.ler_orcamento(centro_float, 2026))),
        "furo: centro 10206911.0 não vira 102069110")
unicode_ = planilha("unicode.xlsx", [["²", "x", "3.1", 0, 0, "0", "Maio", 5, "", None]], ctt=[["²", "A", "", "", "", ""]])
confere('centro de custo "²" inválido' in str(erro_de(lambda: op.ler_orcamento(unicode_, 2026))), "furo: dígito Unicode no centro recusado")
sublinhado = planilha("sublinhado.xlsx", [["010206911", "x", "3.1", 0, 0, "0", "Maio", "1_000", "", None]])
confere("não é número" in str(erro_de(lambda: op.ler_orcamento(sublinhado, 2026))), "furo: '1_000' recusado")
repetido = planilha("repetido.xlsx", [["010206911", "x", "3.1", 0, 0, "0", "Maio", 5, "", None]],
                    ctt=[["010206911", "A", "", "", "", ""], ["01.02.069.11", "B", "", "", "", ""]])
confere("repetido na aba CTT" in str(erro_de(lambda: op.ler_orcamento(repetido, 2026))), "furo: centro repetido na CTT recusado")
confere("Não foi possível ler" in str(erro_de(lambda: op.ler_orcamento(TMP / "nao_existe.xlsx", 2026))), "furo: arquivo inexistente vira ValueError")
longo = planilha("longo.xlsx", [["9" * 30000, "x", "3.1", 0, 0, "0", "Maio", 5, "", None]])
confere(len(str(erro_de(lambda: op.ler_orcamento(longo, 2026)))) < 300, "furo: centro gigante não infla a mensagem")

# 4. arquivo e parâmetros
sem_ctt = planilha("sem_ctt.xlsx", [], abas=("DADOS",))
confere("não tem a aba CTT" in str(erro_de(lambda: op.ler_orcamento(sem_ctt, 2026))), "arquivo: aba faltando dá erro claro")
sem_coluna = planilha("sem_coluna.xlsx", [["010206911", "3.1", "0", "Maio", 10]],
                      cab_dados=["CQ2_CCUSTO", "CQ2_CONTA", "CQ2_TPSALD", "MES", "ORCADO"])
confere("REALIZADO" in str(erro_de(lambda: op.ler_orcamento(sem_coluna, 2026))), "arquivo: coluna obrigatória faltando é nomeada")
so_realizado = planilha("so_realizado.xlsx", [["010206911", "x", "3.1", 0, 0, "1", "Maio", "", 5, None]])
confere("nenhuma linha de orçado" in str(erro_de(lambda: op.ler_orcamento(so_realizado, 2026))), "arquivo: sem orçado dá erro")
for ano in (1999, 2101, "2026", True):
    confere("Ano" in str(erro_de(lambda: op.ler_orcamento(boa, ano))), f"parâmetro: ano {ano!r} recusado")
confere("não é uma planilha" in str(erro_de(lambda: op.conferir_xlsx(b"MZ\x90\x00 exe renomeado"))), "arquivo: .exe renomeado recusado")
confere(erro_de(lambda: op.conferir_xlsx(boa.read_bytes())) is None, "arquivo: xlsx de verdade passa na conferência")

print()
print(f"{len(falhas)} falha(s)" if falhas else "prova_orcamento_leitor: tudo ok")
sys.exit(1 if falhas else 0)
